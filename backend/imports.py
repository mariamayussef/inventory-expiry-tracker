import csv, io, json, zipfile
from datetime import datetime, date
from decimal import Decimal
from fastapi import HTTPException
from openpyxl import load_workbook, Workbook
from psycopg.types.json import Jsonb
from .domain import *

PRODUCT_COLUMNS=['sku','name','category','barcode','unit','whole_units','expiry_tracked','threshold']
STOCK_COLUMNS=['sku','lot','expiry_date','manufacture_date','quantity','supplier','location','reference','reason']

def read_upload(content, filename, mapping, date_format, kind):
    if len(content)>2*1024*1024: fail('Maximum upload size is 2 MB',413)
    if date_format not in ('ISO','DD/MM/YYYY','MM/DD/YYYY'): fail('Select an explicit supported date format')
    if kind not in ('products','opening'): fail('Import kind must be products or opening')
    try:
        if filename.lower().endswith('.csv'):
            rows=list(csv.DictReader(io.StringIO(content.decode('utf-8-sig'))))
        elif filename.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist())>12*1024*1024: fail('Workbook expands beyond 12 MB')
            wb=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
            sheet=wb.active
            if sheet.max_row and sheet.max_row>501: fail('Maximum 500 rows')
            values=list(sheet.values); wb.close()
            rows=[dict(zip([str(x) for x in values[0]],r)) for r in values[1:]] if values else []
        else: fail('Only CSV and XLSX uploads are supported')
    except HTTPException: raise
    except Exception: fail('Could not read this file; check encoding and workbook format')
    if not rows or len(rows)>500: fail('Upload between 1 and 500 data rows')
    expected=PRODUCT_COLUMNS if kind=='products' else STOCK_COLUMNS
    result=[]
    for raw in rows:
        row={col:raw.get(mapping.get(col,col),'') for col in expected}
        row={k:(v.strip() if isinstance(v,str) else v) for k,v in row.items()}
        for k in ('expiry_date','manufacture_date'):
            if row.get(k):
                v=row[k]
                if isinstance(v,(date,datetime)): row[k]=v.date().isoformat() if isinstance(v,datetime) else v.isoformat()
                elif date_format!='ISO':
                    try: row[k]=datetime.strptime(str(v),'%d/%m/%Y' if date_format=='DD/MM/YYYY' else '%m/%d/%Y').date().isoformat()
                    except ValueError: row[k]='INVALID_DATE'
        if kind=='products':
            for k in ('whole_units','expiry_tracked'):
                v=str(row[k]).lower()
                row[k]=True if v in ('true','1','yes','') else False if v in ('false','0','no') else v
            row['threshold']=str(row['threshold'] or '10')
            row['unit']=row['unit'] or 'each'; row['category']=row['category'] or 'General'
        else:
            row['quantity']=str(row['quantity']); row['reason']=row['reason'] or 'Opening stock import'
        result.append(clean(row))
    return result

def apply_row(c,kind,row,actor,import_id=None):
    if kind=='products': return create_product(c,row,actor)
    p=product_by_sku(c,row['sku'])
    return command(c,dict(row,type='opening',product_id=str(p['id'])),actor,import_id)

def stage(c,kind,rows,actor):
    errors=[]; seen=set()
    # The outer rollback discards all preview writes; inner savepoints collect row errors.
    c.execute('SAVEPOINT preview_all')
    try:
        for i,row in enumerate(rows,2):
            identity=row.get('sku') if kind=='products' else (row.get('sku'),row.get('lot'))
            if identity in seen:
                errors.append({'row':i,'message':'Duplicate SKU or SKU/lot within upload'}); continue
            seen.add(identity)
            c.execute('SAVEPOINT preview_row')
            try: apply_row(c,kind,row,actor)
            except HTTPException as e:
                c.execute('ROLLBACK TO SAVEPOINT preview_row'); errors.append({'row':i,'message':str(e.detail)})
            except Exception as e:
                c.execute('ROLLBACK TO SAVEPOINT preview_row'); errors.append({'row':i,'message':'Invalid row or database constraint: '+type(e).__name__})
            c.execute('RELEASE SAVEPOINT preview_row')
    finally:
        c.execute('ROLLBACK TO SAVEPOINT preview_all'); c.execute('RELEASE SAVEPOINT preview_all')
    ident=uid()
    c.execute('INSERT INTO imports(id,kind,rows,errors,actor) VALUES(%s,%s,%s,%s,%s)',(ident,kind,Jsonb(rows),Jsonb(errors),actor['id']))
    return {'id':ident,'kind':kind,'rows':rows,'errors':errors,'can_commit':not errors}

def commit_import(c,ident,actor):
    staged=c.execute('SELECT * FROM imports WHERE id=%s FOR UPDATE',(ident,)).fetchone()
    if not staged: fail('Import not found',404)
    if staged['errors']: fail('Fix all validation errors and preview a new file')
    if staged['committed_at']: return {'import_id':ident,'rows_posted':len(staged['rows']),'already_committed':True}
    # Deterministic product lock order for stock imports avoids lock inversion.
    if staged['kind']=='opening':
        skus=sorted({r['sku'] for r in staged['rows']})
        c.execute('SELECT id FROM products WHERE sku=ANY(%s) ORDER BY id FOR UPDATE',(skus,)).fetchall()
    for row in staged['rows']: apply_row(c,staged['kind'],row,actor,ident)
    c.execute('UPDATE imports SET committed_at=now() WHERE id=%s',(ident,))
    audit(c,actor,'import_committed',ident,{'rows':len(staged['rows'])})
    return {'import_id':ident,'rows_posted':len(staged['rows'])}

def safe_cell(v):
    text=str(v) if v is not None else ''
    return "'"+text if text.lstrip().startswith(('=','+','-','@','\t','\r')) else text

def spreadsheet(rows,format):
    columns=list(rows[0]) if rows else ['no_results']
    if format=='xlsx':
        wb=Workbook(); ws=wb.active; ws.title='Inventory export'; ws.append(columns)
        for row in rows: ws.append([safe_cell(serial(v) if isinstance(v,(datetime,date,Decimal)) else v) for v in row.values()])
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(44,max(16,len(str(col[0].value))+3))
        out=io.BytesIO(); wb.save(out); return out.getvalue()
    if format!='csv': fail('Export format must be csv or xlsx')
    out=io.StringIO(); writer=csv.writer(out); writer.writerow(columns)
    for row in rows: writer.writerow([safe_cell(serial(v) if isinstance(v,(datetime,date,Decimal)) else v) for v in row.values()])
    return out.getvalue().encode('utf-8-sig')
